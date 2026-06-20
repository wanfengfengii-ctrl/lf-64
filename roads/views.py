from django.shortcuts import render, get_object_or_404, redirect
from django.http import JsonResponse, HttpResponse
from django.urls import reverse
from django.db.models import Count, Avg, Q, Max, Min, F, Exists, OuterRef
from django.core.exceptions import ValidationError
from django.utils import timezone
from django.contrib import messages
from datetime import datetime, date, timedelta
from collections import defaultdict
import csv
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from .models import (
    RoadSection, Point, InspectionRecord, Photo, Alert, TaskOrder,
    Hazard, HazardDisposal, RoadPassageStatus,
    OpenSchedule, TemporaryControl, VisitorFlowRecord,
    WEAR_LEVEL_CHOICES, POINT_TYPE_CHOICES, ROAD_STATUS_CHOICES,
    ALERT_LEVEL_CHOICES, ALERT_TYPE_CHOICES, TASK_STATUS_CHOICES,
    HAZARD_LOCATION_TYPE_CHOICES, HAZARD_TYPE_CHOICES,
    HAZARD_LEVEL_CHOICES, HAZARD_STATUS_CHOICES,
    CONTROL_SUGGESTION_CHOICES, PASSAGE_STATUS_CHOICES,
    DISPOSAL_TYPE_CHOICES, SEASON_CHOICES, WEATHER_CONDITION_CHOICES,
    CONTROL_TYPE_CHOICES, FLOW_RECORD_TYPE_CHOICES,
    OPEN_STATUS_CHOICES, MAINTENANCE_STATUS_CHOICES,
)
from .forms import (
    RoadSectionForm, PointForm, InspectionRecordForm,
    PhotoForm, TaskOrderForm, TaskDispatchForm, TaskRectifyForm,
    TaskReviewForm, AlertForm, DataExportForm,
    HazardForm, HazardAssessForm, HazardStatusForm,
    HazardDisposalForm, RoadPassageStatusForm,
    OpenScheduleForm, TemporaryControlForm, VisitorFlowRecordForm,
)


def get_high_risk_annotation():
    return Exists(
        InspectionRecord.objects.filter(
            point_id=OuterRef('pk'),
            handled=False,
            wear_level__gte=3,
        )
    )


def annotate_high_risk(queryset):
    return queryset.annotate(
        is_high_risk_annotated=get_high_risk_annotation(),
        has_unhandled_critical_annotated=Exists(
            InspectionRecord.objects.filter(
                point_id=OuterRef('pk'),
                handled=False,
                wear_level=4,
            )
        ),
        worst_unhandled_wear_annotated=Max(
            'inspections__wear_level',
            filter=Q(inspections__handled=False)
        ),
    )


def dashboard(request):
    total_roads = RoadSection.objects.count()
    total_points = Point.objects.count()
    total_inspections = InspectionRecord.objects.count()
    high_risk_points = annotate_high_risk(Point.objects).filter(
        is_high_risk_annotated=True
    ).count()
    critical_points = annotate_high_risk(Point.objects).filter(
        has_unhandled_critical_annotated=True
    )

    unread_alerts = Alert.objects.filter(is_read=False, is_resolved=False).count()
    critical_alerts = Alert.objects.filter(alert_level='critical', is_resolved=False).count()
    pending_tasks = TaskOrder.objects.filter(status='pending').count()
    rectifying_tasks = TaskOrder.objects.filter(status='rectifying').count()
    overdue_tasks = TaskOrder.objects.exclude(status='closed').filter(deadline__lt=timezone.now().date()).count()

    recent_inspections = InspectionRecord.objects.select_related('point', 'point__road_section').order_by('-inspection_date')[:10]
    recent_alerts = Alert.objects.select_related('point', 'point__road_section').order_by('-created_at')[:8]
    recent_tasks = TaskOrder.objects.select_related('point', 'point__road_section').order_by('-created_at')[:8]

    roads_with_issues = RoadSection.objects.annotate(
        unhandled_critical=Count('points', filter=Q(
            points__inspections__wear_level=4,
            points__inspections__handled=False
        ), distinct=True)
    ).filter(unhandled_critical__gt=0).order_by('-unhandled_critical')

    wear_distribution = InspectionRecord.objects.values('wear_level').annotate(
        count=Count('id')
    ).order_by('wear_level')

    wear_stats_dict = {level: 0 for level, _ in WEAR_LEVEL_CHOICES}
    for item in wear_distribution:
        wear_stats_dict[item['wear_level']] = item['count']

    wear_stats_list = []
    for level, label in WEAR_LEVEL_CHOICES:
        count = wear_stats_dict.get(level, 0)
        percentage = round(count * 100 / total_inspections, 1) if total_inspections > 0 else 0
        wear_stats_list.append({
            'level': level,
            'label': label,
            'count': count,
            'percentage': percentage,
        })

    task_status_stats = TaskOrder.objects.values('status').annotate(count=Count('id'))
    task_stats_dict = {s: 0 for s, _ in TASK_STATUS_CHOICES}
    for item in task_status_stats:
        task_stats_dict[item['status']] = item['count']

    total_hazards = Hazard.objects.count()
    active_hazards = Hazard.objects.exclude(status__in=['resolved', 'closed'])
    critical_hazards = active_hazards.filter(hazard_level='critical').count()
    severe_hazards = active_hazards.filter(hazard_level='severe').count()
    warning_hazards = active_hazards.filter(hazard_level='warning').count()
    disposing_hazards = active_hazards.filter(status='disposing').count()
    overdue_hazards = active_hazards.filter(disposal_deadline__lt=timezone.now().date()).count()
    roads_with_hazards = RoadSection.objects.annotate(
        hazard_count=Count('hazards', filter=~Q(hazards__status__in=['resolved', 'closed']), distinct=True)
    ).filter(hazard_count__gt=0).order_by('-hazard_count')[:5]

    recent_hazards = Hazard.objects.select_related(
        'road_section', 'point'
    ).prefetch_related('disposals').order_by('-reported_at')[:8]

    hazard_level_distribution = active_hazards.values('hazard_level').annotate(
        count=Count('id')
    ).order_by('hazard_level')
    hazard_level_dict = {level: 0 for level, _ in HAZARD_LEVEL_CHOICES if level != 'safe'}
    for item in hazard_level_distribution:
        hazard_level_dict[item['hazard_level']] = item['count']

    passage_status_stats = RoadPassageStatus.objects.values('passage_status').annotate(
        count=Count('id')
    )
    passage_stats_dict = {s: 0 for s, _ in PASSAGE_STATUS_CHOICES}
    for item in passage_status_stats:
        passage_stats_dict[item['passage_status']] = item['count']

    active_schedules = OpenSchedule.objects.filter(is_active=True).count()
    active_controls = TemporaryControl.objects.filter(
        is_active=True, end_time__gte=timezone.now()
    ).count()
    today_flows = VisitorFlowRecord.objects.filter(
        record_date=timezone.now().date()
    )

    return render(request, 'roads/dashboard.html', {
        'total_roads': total_roads,
        'total_points': total_points,
        'total_inspections': total_inspections,
        'high_risk_points': high_risk_points,
        'critical_points': critical_points,
        'unread_alerts': unread_alerts,
        'critical_alerts': critical_alerts,
        'pending_tasks': pending_tasks,
        'rectifying_tasks': rectifying_tasks,
        'overdue_tasks': overdue_tasks,
        'recent_inspections': recent_inspections,
        'recent_alerts': recent_alerts,
        'recent_tasks': recent_tasks,
        'roads_with_issues': roads_with_issues,
        'wear_stats_list': wear_stats_list,
        'task_stats_dict': task_stats_dict,
        'task_status_choices': dict(TASK_STATUS_CHOICES),
        'total_hazards': total_hazards,
        'active_hazards_count': active_hazards.count(),
        'critical_hazards': critical_hazards,
        'severe_hazards': severe_hazards,
        'warning_hazards': warning_hazards,
        'disposing_hazards': disposing_hazards,
        'overdue_hazards': overdue_hazards,
        'recent_hazards': recent_hazards,
        'hazard_level_dict': hazard_level_dict,
        'hazard_level_choices': dict(HAZARD_LEVEL_CHOICES),
        'hazard_status_choices': dict(HAZARD_STATUS_CHOICES),
        'roads_with_hazards': roads_with_hazards,
        'passage_stats_dict': passage_stats_dict,
        'passage_status_choices': dict(PASSAGE_STATUS_CHOICES),
        'active_schedules': active_schedules,
        'active_controls': active_controls,
        'today_flow_count': today_flows.count(),
    })


def map_view(request):
    high_risk = request.GET.get('high_risk', '0') == '1'
    date_param = request.GET.get('date', '')
    points = Point.objects.select_related('road_section').all()
    if high_risk:
        points = points.filter(
            inspections__wear_level__gte=3,
            inspections__handled=False
        ).distinct()

    return render(request, 'roads/map.html', {
        'points': points,
        'high_risk_filter': high_risk,
        'date': date_param,
        'wear_level_choices': dict(WEAR_LEVEL_CHOICES),
        'point_type_choices': dict(POINT_TYPE_CHOICES),
    })


def charts_view(request):
    roads = RoadSection.objects.all()
    selected_road_id = request.GET.get('road_id')
    months_count = int(request.GET.get('months', '6'))
    selected_point_id = request.GET.get('point')
    points = Point.objects.select_related('road_section').all().order_by('code')

    today = timezone.now().date()
    labels = []
    for i in range(months_count - 1, -1, -1):
        month_date = date(today.year, today.month, 1)
        if today.month - i <= 0:
            year = today.year - 1
            month = today.month - i + 12
        else:
            year = today.year
            month = today.month - i
        labels.append(f'{year}年{month}月')

    road_data = []
    if selected_road_id:
        road = get_object_or_404(RoadSection, pk=selected_road_id)
        monthly_avg = []
        for i in range(months_count - 1, -1, -1):
            if today.month - i <= 0:
                year = today.year - 1
                month = today.month - i + 12
            else:
                year = today.year
                month = today.month - i
            inspections = InspectionRecord.objects.filter(
                point__road_section=road,
                inspection_date__year=year,
                inspection_date__month=month
            )
            if inspections.exists():
                avg = inspections.aggregate(avg=Avg('wear_level'))['avg']
                monthly_avg.append(round(avg, 2))
            else:
                monthly_avg.append(0)
        road_data.append({
            'name': road.name,
            'data': monthly_avg,
        })
    else:
        for road in roads:
            monthly_avg = []
            for i in range(months_count - 1, -1, -1):
                if today.month - i <= 0:
                    year = today.year - 1
                    month = today.month - i + 12
                else:
                    year = today.year
                    month = today.month - i
                inspections = InspectionRecord.objects.filter(
                    point__road_section=road,
                    inspection_date__year=year,
                    inspection_date__month=month
                )
                if inspections.exists():
                    avg = inspections.aggregate(avg=Avg('wear_level'))['avg']
                    monthly_avg.append(round(avg, 2))
                else:
                    monthly_avg.append(0)
            road_data.append({
                'name': road.name,
                'data': monthly_avg,
            })

    wear_count_data = defaultdict(lambda: [0] * len(labels))
    for i in range(months_count - 1, -1, -1):
        idx = months_count - 1 - i
        if today.month - i <= 0:
            year = today.year - 1
            month = today.month - i + 12
        else:
            year = today.year
            month = today.month - i
        for level, _ in WEAR_LEVEL_CHOICES:
            count = InspectionRecord.objects.filter(
                inspection_date__year=year,
                inspection_date__month=month,
                wear_level=level
            ).count()
            wear_count_data[level][idx] = count

    hazard_level_count = defaultdict(lambda: [0] * len(labels))
    hazard_total = [0] * len(labels)
    hazard_critical_severe = [0] * len(labels)

    for i in range(months_count - 1, -1, -1):
        idx = months_count - 1 - i
        if today.month - i <= 0:
            year = today.year - 1
            month = today.month - i + 12
        else:
            year = today.year
            month = today.month - i

        month_hazards = Hazard.objects.filter(
            reported_date__year=year,
            reported_date__month=month
        )
        total = month_hazards.count()
        hazard_total[idx] = total

        critical_severe = month_hazards.filter(
            hazard_level__in=['critical', 'severe']
        ).count()
        hazard_critical_severe[idx] = critical_severe

        for level, _ in HAZARD_LEVEL_CHOICES:
            if level == 'safe':
                continue
            count = month_hazards.filter(hazard_level=level).count()
            hazard_level_count[level][idx] = count

    return render(request, 'roads/charts.html', {
        'roads': roads,
        'points': points,
        'selected_road_id': int(selected_road_id) if selected_road_id else None,
        'selected_point_id': int(selected_point_id) if selected_point_id else None,
        'months_count': months_count,
        'labels': labels,
        'road_data': road_data,
        'wear_level_choices': dict(WEAR_LEVEL_CHOICES),
        'wear_count_data': dict(wear_count_data),
        'hazard_level_count': dict(hazard_level_count),
        'hazard_total': hazard_total,
        'hazard_critical_severe': hazard_critical_severe,
        'hazard_level_choices': {k: v for k, v in HAZARD_LEVEL_CHOICES if k != 'safe'},
    })


def api_wear_data(request):
    road_id = request.GET.get('road_id')
    months_count = int(request.GET.get('months', '6'))
    today = timezone.now().date()

    labels = []
    road_data = {}

    roads = RoadSection.objects.all()
    if road_id:
        roads = roads.filter(pk=road_id)

    for road in roads:
        road_data[road.name] = []

    for i in range(months_count - 1, -1, -1):
        if today.month - i <= 0:
            year = today.year - 1
            month = today.month - i + 12
        else:
            year = today.year
            month = today.month - i
        labels.append(f'{year}年{month}月')

        for road in roads:
            inspections = InspectionRecord.objects.filter(
                point__road_section=road,
                inspection_date__year=year,
                inspection_date__month=month
            )
            if inspections.exists():
                avg = inspections.aggregate(avg=Avg('wear_level'))['avg']
                road_data[road.name].append(round(avg, 2))
            else:
                road_data[road.name].append(0)

    return JsonResponse({
        'labels': labels,
        'datasets': road_data,
    })


def api_points_geo(request):
    high_risk = request.GET.get('high_risk', '0') == '1'
    points = annotate_high_risk(Point.objects.select_related('road_section').all())
    if high_risk:
        points = points.filter(is_high_risk_annotated=True)

    features = []
    for point in points:
        worst_wear = point.worst_unhandled_wear_annotated
        latest = point.get_latest_inspection()
        is_high = point.is_high_risk_annotated
        wear_level = worst_wear if worst_wear else None

        features.append({
            'type': 'Feature',
            'geometry': {
                'type': 'Point',
                'coordinates': [point.longitude, point.latitude],
            },
            'properties': {
                'id': point.id,
                'code': point.code,
                'name': point.name,
                'point_type': point.point_type,
                'point_type_display': point.get_point_type_display(),
                'road_section': point.road_section.name,
                'road_section_id': point.road_section.id,
                'wear_level': wear_level,
                'wear_level_display': dict(WEAR_LEVEL_CHOICES).get(wear_level, '无未处理记录'),
                'is_high_risk': bool(is_high),
                'handled': wear_level is None,
                'has_latest': latest is not None,
                'latest_wear_level': latest.wear_level if latest else None,
                'latest_wear_display': dict(WEAR_LEVEL_CHOICES).get(latest.wear_level, '无记录') if latest else '无记录',
            }
        })

    return JsonResponse({
        'type': 'FeatureCollection',
        'features': features,
    })


def road_list(request):
    roads = RoadSection.objects.annotate(
        point_count=Count('points', distinct=True),
        unhandled_critical=Count('points', filter=Q(
            points__inspections__wear_level=4,
            points__inspections__handled=False
        ), distinct=True)
    ).all().order_by('code')
    return render(request, 'roads/road_list.html', {'roads': roads})


def road_create(request):
    if request.method == 'POST':
        form = RoadSectionForm(request.POST)
        if form.is_valid():
            road = form.save()
            return redirect('roads:road_detail', pk=road.pk)
    else:
        form = RoadSectionForm()
    return render(request, 'roads/road_form.html', {'form': form, 'mode': 'create'})


def road_detail(request, pk):
    road = get_object_or_404(RoadSection.objects.prefetch_related('points', 'points__inspections'), pk=pk)
    points = annotate_high_risk(road.points.all()).order_by('code')
    return render(request, 'roads/road_detail.html', {'road': road, 'points': points})


def road_edit(request, pk):
    road = get_object_or_404(RoadSection, pk=pk)
    if request.method == 'POST':
        form = RoadSectionForm(request.POST, instance=road)
        if form.is_valid():
            form.save()
            return redirect('roads:road_detail', pk=road.pk)
    else:
        form = RoadSectionForm(instance=road)
    return render(request, 'roads/road_form.html', {'form': form, 'road': road, 'mode': 'edit'})


def road_delete(request, pk):
    road = get_object_or_404(RoadSection, pk=pk)
    if request.method == 'POST':
        road.delete()
        return redirect('roads:road_list')
    return render(request, 'roads/road_confirm_delete.html', {'road': road})


def point_list(request):
    road_filter = request.GET.get('road')
    type_filter = request.GET.get('type')
    risk_filter = request.GET.get('risk')

    points = annotate_high_risk(
        Point.objects.select_related('road_section').all()
    )

    if road_filter:
        points = points.filter(road_section_id=road_filter)
    if type_filter:
        points = points.filter(point_type=type_filter)
    if risk_filter == 'high':
        points = points.filter(is_high_risk_annotated=True)

    points = points.annotate(
        latest_wear=Avg('inspections__wear_level')
    ).order_by('code')

    roads = RoadSection.objects.all()
    return render(request, 'roads/point_list.html', {
        'points': points,
        'roads': roads,
        'road_filter': road_filter,
        'type_filter': type_filter,
        'risk_filter': risk_filter,
        'point_type_choices': dict(POINT_TYPE_CHOICES),
        'wear_level_choices': dict(WEAR_LEVEL_CHOICES),
    })


def point_create(request):
    if request.method == 'POST':
        form = PointForm(request.POST)
        if form.is_valid():
            try:
                point = form.save()
                messages.success(request, f'点位 {point.code} 创建成功。')
                return redirect('roads:point_detail', pk=point.pk)
            except Exception as e:
                form.add_error(None, f'保存失败：{e}')
    else:
        form = PointForm()
    return render(request, 'roads/point_form.html', {'form': form, 'mode': 'create'})


def point_detail(request, pk):
    point = get_object_or_404(annotate_high_risk(
        Point.objects.select_related('road_section').prefetch_related(
            'inspections', 'photos', 'task_orders', 'alerts'
        )
    ), pk=pk)
    inspections = point.inspections.all().order_by('-inspection_date')
    photos = point.photos.all().order_by('-uploaded_at')[:12]
    task_orders = point.task_orders.all().order_by('-created_at')
    alerts = point.alerts.all().order_by('-created_at')[:10]
    return render(request, 'roads/point_detail.html', {
        'point': point,
        'inspections': inspections,
        'photos': photos,
        'task_orders': task_orders,
        'alerts': alerts,
        'wear_level_choices': dict(WEAR_LEVEL_CHOICES),
        'task_status_choices': dict(TASK_STATUS_CHOICES),
    })


def point_edit(request, pk):
    point = get_object_or_404(Point, pk=pk)
    if request.method == 'POST':
        form = PointForm(request.POST, instance=point)
        if form.is_valid():
            try:
                form.save()
                messages.success(request, f'点位 {point.code} 更新成功。')
                return redirect('roads:point_detail', pk=point.pk)
            except Exception as e:
                form.add_error(None, f'保存失败：{e}')
    else:
        form = PointForm(instance=point)
    return render(request, 'roads/point_form.html', {'form': form, 'point': point, 'mode': 'edit'})


def point_delete(request, pk):
    point = get_object_or_404(Point, pk=pk)
    road_pk = point.road_section.pk
    if request.method == 'POST':
        point.delete()
        return redirect('roads:road_detail', pk=road_pk)
    return render(request, 'roads/point_confirm_delete.html', {'point': point})


def inspection_list(request):
    point_filter = request.GET.get('point')
    level_filter = request.GET.get('level')
    handled_filter = request.GET.get('handled')

    inspections = InspectionRecord.objects.select_related('point', 'point__road_section').all()

    if point_filter:
        inspections = inspections.filter(point_id=point_filter)
    if level_filter:
        inspections = inspections.filter(wear_level=int(level_filter))
    if handled_filter == 'yes':
        inspections = inspections.filter(handled=True)
    elif handled_filter == 'no':
        inspections = inspections.filter(handled=False)

    inspections = inspections.order_by('-inspection_date')
    points = Point.objects.all().order_by('code')

    return render(request, 'roads/inspection_list.html', {
        'inspections': inspections,
        'points': points,
        'point_filter': point_filter,
        'level_filter': level_filter,
        'handled_filter': handled_filter,
        'wear_level_choices': dict(WEAR_LEVEL_CHOICES),
    })


def inspection_create(request):
    if request.method == 'POST':
        form = InspectionRecordForm(request.POST)
        if form.is_valid():
            try:
                insp = form.save(commit=False)
                insp.full_clean()
                insp = form.save()
                return redirect('roads:inspection_detail', pk=insp.pk)
            except ValidationError as e:
                for field, errors in e.message_dict.items():
                    for error in errors:
                        form.add_error(field, error)
    else:
        form = InspectionRecordForm(initial={'inspection_date': timezone.now().date()})
    return render(request, 'roads/inspection_form.html', {'form': form, 'mode': 'create'})


def inspection_detail(request, pk):
    inspection = get_object_or_404(
        InspectionRecord.objects.select_related('point', 'point__road_section'),
        pk=pk
    )
    return render(request, 'roads/inspection_detail.html', {
        'inspection': inspection,
        'wear_level_choices': dict(WEAR_LEVEL_CHOICES),
    })


def inspection_edit(request, pk):
    inspection = get_object_or_404(InspectionRecord, pk=pk)
    if request.method == 'POST':
        form = InspectionRecordForm(request.POST, instance=inspection)
        if form.is_valid():
            try:
                insp = form.save(commit=False)
                insp.full_clean()
                form.save()
                return redirect('roads:inspection_detail', pk=inspection.pk)
            except ValidationError as e:
                for field, errors in e.message_dict.items():
                    for error in errors:
                        form.add_error(field, error)
    else:
        form = InspectionRecordForm(instance=inspection)
    return render(request, 'roads/inspection_form.html', {'form': form, 'inspection': inspection, 'mode': 'edit'})


def inspection_delete(request, pk):
    inspection = get_object_or_404(InspectionRecord, pk=pk)
    point_pk = inspection.point.pk
    if request.method == 'POST':
        inspection.delete()
        return redirect('roads:point_detail', pk=point_pk)
    return render(request, 'roads/inspection_confirm_delete.html', {'inspection': inspection})


# ==================== 风险预警模块 ====================

def alert_list(request):
    level_filter = request.GET.get('level')
    type_filter = request.GET.get('type')
    status_filter = request.GET.get('status')
    road_filter = request.GET.get('road')

    alerts = Alert.objects.select_related('point', 'point__road_section', 'related_inspection').all()

    if level_filter:
        alerts = alerts.filter(alert_level=level_filter)
    if type_filter:
        alerts = alerts.filter(alert_type=type_filter)
    if status_filter == 'unread':
        alerts = alerts.filter(is_read=False)
    elif status_filter == 'unresolved':
        alerts = alerts.filter(is_resolved=False)
    elif status_filter == 'resolved':
        alerts = alerts.filter(is_resolved=True)
    if road_filter:
        alerts = alerts.filter(point__road_section_id=road_filter)

    alerts = alerts.order_by('-created_at')
    roads = RoadSection.objects.all()

    return render(request, 'roads/alert_list.html', {
        'alerts': alerts,
        'roads': roads,
        'level_filter': level_filter,
        'type_filter': type_filter,
        'status_filter': status_filter,
        'road_filter': road_filter,
        'alert_level_choices': dict(ALERT_LEVEL_CHOICES),
        'alert_type_choices': dict(ALERT_TYPE_CHOICES),
    })


def alert_detail(request, pk):
    alert = get_object_or_404(
        Alert.objects.select_related('point', 'point__road_section', 'related_inspection'),
        pk=pk
    )
    if not alert.is_read:
        alert.is_read = True
        alert.save()
    return render(request, 'roads/alert_detail.html', {
        'alert': alert,
        'alert_level_choices': dict(ALERT_LEVEL_CHOICES),
        'alert_type_choices': dict(ALERT_TYPE_CHOICES),
    })


def alert_mark_read(request, pk):
    alert = get_object_or_404(Alert, pk=pk)
    alert.is_read = True
    alert.save()
    messages.success(request, '预警已标记为已读。')
    return redirect(request.META.get('HTTP_REFERER', 'roads:alert_list'))


def alert_resolve(request, pk):
    alert = get_object_or_404(Alert, pk=pk)
    alert.resolve()
    messages.success(request, '预警已处理完成。')
    return redirect(request.META.get('HTTP_REFERER', 'roads:alert_list'))


def alert_create(request):
    if request.method == 'POST':
        form = AlertForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, '预警创建成功。')
            return redirect('roads:alert_list')
    else:
        form = AlertForm()
    return render(request, 'roads/alert_form.html', {'form': form, 'mode': 'create'})


def alert_generate_auto(request):
    count = 0
    today = timezone.now().date()

    critical_inspections = InspectionRecord.objects.filter(
        wear_level=4, handled=False
    ).select_related('point')
    for insp in critical_inspections:
        existing = Alert.objects.filter(
            point=insp.point,
            alert_type='wear_critical',
            related_inspection=insp,
            is_resolved=False
        ).exists()
        if not existing:
            Alert.objects.create(
                point=insp.point,
                alert_type='wear_critical',
                alert_level='critical',
                message=f'点位[{insp.point.code}] {insp.point.name} 发现严重磨损（等级4），请立即处理。巡查日期：{insp.inspection_date}，巡查员：{insp.inspector}',
                related_inspection=insp,
            )
            count += 1

    points = Point.objects.all()
    for point in points:
        inspections = list(point.inspections.order_by('inspection_date'))
        if len(inspections) >= 2:
            latest = inspections[-1]
            prev = inspections[-2]
            if latest.wear_level > prev.wear_level and not latest.handled:
                existing = Alert.objects.filter(
                    point=point,
                    alert_type='wear_worsening',
                    related_inspection=latest,
                    is_resolved=False
                ).exists()
                if not existing:
                    Alert.objects.create(
                        point=point,
                        alert_type='wear_worsening',
                        alert_level='warning',
                        message=f'点位[{point.code}] {point.name} 磨损状况恶化：从{prev.get_wear_level_display()}升级为{latest.get_wear_level_display()}。巡查日期：{latest.inspection_date}',
                        related_inspection=latest,
                    )
                    count += 1

    overdue_tasks = TaskOrder.objects.exclude(status='closed').filter(deadline__lt=today)
    for task in overdue_tasks:
        existing = Alert.objects.filter(
            point=task.point,
            alert_type='task_overdue',
            is_resolved=False,
            created_at__date=today
        ).exists()
        if not existing:
            Alert.objects.create(
                point=task.point,
                alert_type='task_overdue',
                alert_level='warning',
                message=f'工单[{task.title}]已超期。指派人：{task.assigned_to or "未指派"}，截止日期：{task.deadline}，当前状态：{task.get_status_display()}',
            )
            count += 1

    thirty_days_ago = today - timedelta(days=30)
    long_unhandled = InspectionRecord.objects.filter(
        handled=False,
        wear_level__gte=3,
        inspection_date__lt=thirty_days_ago
    ).select_related('point')
    for insp in long_unhandled:
        existing = Alert.objects.filter(
            point=insp.point,
            alert_type='unhandled_long',
            related_inspection=insp,
            is_resolved=False
        ).exists()
        if not existing:
            Alert.objects.create(
                point=insp.point,
                alert_type='unhandled_long',
                alert_level='critical' if insp.wear_level == 4 else 'warning',
                message=f'点位[{insp.point.code}] {insp.point.name} 存在{insp.get_wear_level_display()}磨损超过30天未处理。巡查日期：{insp.inspection_date}',
                related_inspection=insp,
            )
            count += 1

    messages.success(request, f'自动生成预警完成，共生成 {count} 条预警。')
    return redirect('roads:alert_list')


# ==================== 整改闭环 / 工单模块 ====================

def task_list(request):
    status_filter = request.GET.get('status')
    priority_filter = request.GET.get('priority')
    road_filter = request.GET.get('road')
    assigned_filter = request.GET.get('assigned')
    overdue_filter = request.GET.get('overdue')

    tasks = TaskOrder.objects.select_related('point', 'point__road_section', 'inspection').all()

    if status_filter:
        tasks = tasks.filter(status=status_filter)
    if priority_filter:
        tasks = tasks.filter(priority=priority_filter)
    if road_filter:
        tasks = tasks.filter(point__road_section_id=road_filter)
    if assigned_filter == 'yes':
        tasks = tasks.exclude(assigned_to='')
    elif assigned_filter == 'no':
        tasks = tasks.filter(assigned_to='')
    if overdue_filter == '1':
        today = timezone.now().date()
        tasks = tasks.exclude(status='closed').filter(deadline__lt=today)

    tasks = tasks.order_by('-created_at')
    roads = RoadSection.objects.all()

    stats = {
        'total': tasks.count(),
        'pending': tasks.filter(status='pending').count(),
        'dispatched': tasks.filter(status='dispatched').count(),
        'rectifying': tasks.filter(status='rectifying').count(),
        'pending_review': tasks.filter(status='pending_review').count(),
        'closed': tasks.filter(status='closed').count(),
    }

    return render(request, 'roads/task_list.html', {
        'tasks': tasks,
        'roads': roads,
        'status_filter': status_filter,
        'priority_filter': priority_filter,
        'road_filter': road_filter,
        'assigned_filter': assigned_filter,
        'overdue_filter': overdue_filter,
        'task_status_choices': dict(TASK_STATUS_CHOICES),
        'stats': stats,
    })


def task_detail(request, pk):
    task = get_object_or_404(
        TaskOrder.objects.select_related('point', 'point__road_section', 'inspection'),
        pk=pk
    )
    photos = Photo.objects.filter(
        point=task.point,
        photo_type__in=['rectification', 'verification']
    ).order_by('-uploaded_at')
    return render(request, 'roads/task_detail.html', {
        'task': task,
        'photos': photos,
        'task_status_choices': dict(TASK_STATUS_CHOICES),
    })


def task_create(request, inspection_pk=None):
    initial = {}
    if inspection_pk:
        inspection = get_object_or_404(InspectionRecord, pk=inspection_pk)
        initial['inspection'] = inspection
        initial['point'] = inspection.point
        initial['title'] = f'{inspection.point.code} - {inspection.get_wear_level_display()}磨损整改'
        initial['description'] = inspection.maintenance_suggestion or inspection.wear_description
        if inspection.wear_level == 4:
            initial['priority'] = 'urgent'
        elif inspection.wear_level == 3:
            initial['priority'] = 'high'

    if request.method == 'POST':
        form = TaskOrderForm(request.POST)
        if form.is_valid():
            task = form.save(commit=False)
            task.status = 'pending'
            task.save()
            messages.success(request, '工单创建成功，等待派发。')
            return redirect('roads:task_detail', pk=task.pk)
    else:
        form = TaskOrderForm(initial=initial)
    return render(request, 'roads/task_form.html', {'form': form, 'mode': 'create'})


def task_dispatch(request, pk):
    task = get_object_or_404(TaskOrder, pk=pk)
    if task.status not in ['pending', 'rejected']:
        messages.error(request, '当前状态不允许派发。')
        return redirect('roads:task_detail', pk=pk)

    if request.method == 'POST':
        form = TaskDispatchForm(request.POST, instance=task)
        if form.is_valid():
            t = form.save(commit=False)
            t.status = 'dispatched'
            t.dispatched_at = timezone.now()
            t.save()
            messages.success(request, f'工单已派发给 {t.assigned_to or "相关人员"}。')
            return redirect('roads:task_detail', pk=task.pk)
    else:
        form = TaskDispatchForm(instance=task, initial={
            'deadline': task.deadline or (timezone.now().date() + timedelta(days=7)),
            'priority': task.priority,
        })
    return render(request, 'roads/task_dispatch.html', {'form': form, 'task': task})


def task_rectify(request, pk):
    task = get_object_or_404(TaskOrder, pk=pk)
    if task.status not in ['dispatched', 'rectifying']:
        messages.error(request, '当前状态不允许提交整改。')
        return redirect('roads:task_detail', pk=pk)

    if request.method == 'POST':
        form = TaskRectifyForm(request.POST, instance=task)
        if form.is_valid():
            t = form.save(commit=False)
            if t.status == 'dispatched':
                t.status = 'rectifying'
            if t.rectification_result and t.rectification_result.strip():
                t.status = 'pending_review'
                t.rectified_at = timezone.now()
            t.save()
            messages.success(request, '整改信息已提交。' if t.status == 'rectifying' else '整改结果已提交，等待复核。')
            return redirect('roads:task_detail', pk=task.pk)
    else:
        form = TaskRectifyForm(instance=task)
    return render(request, 'roads/task_rectify.html', {'form': form, 'task': task})


def task_review(request, pk):
    task = get_object_or_404(TaskOrder, pk=pk)
    if task.status != 'pending_review':
        messages.error(request, '当前状态不允许复核。')
        return redirect('roads:task_detail', pk=pk)

    if request.method == 'POST':
        form = TaskReviewForm(request.POST, instance=task)
        action = request.POST.get('action')
        if form.is_valid():
            t = form.save(commit=False)
            t.reviewed_at = timezone.now()
            if action == 'pass':
                t.status = 'closed'
                t.closed_at = timezone.now()
                if t.inspection and not t.inspection.handled:
                    t.inspection.handled = True
                    t.inspection.handled_date = timezone.now().date()
                    t.inspection.save()
                messages.success(request, '复核通过，工单已闭环。')
            else:
                t.status = 'rejected'
                messages.warning(request, '复核不通过，请重新整改。')
            t.save()
            return redirect('roads:task_detail', pk=task.pk)
    else:
        form = TaskReviewForm(instance=task)
    return render(request, 'roads/task_review.html', {'form': form, 'task': task})


def task_status_progress(request, pk):
    task = get_object_or_404(TaskOrder, pk=pk)
    new_status = request.POST.get('status')
    valid_transitions = {
        'dispatched': 'rectifying',
    }
    if new_status in valid_transitions.values() and task.status == 'dispatched':
        task.status = new_status
        task.save()
        messages.success(request, '工单状态已更新。')
    else:
        messages.error(request, '无效的状态转换。')
    return redirect('roads:task_detail', pk=pk)


# ==================== 照片档案模块 ====================

def photo_list(request):
    point_filter = request.GET.get('point')
    type_filter = request.GET.get('type')
    road_filter = request.GET.get('road')
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')

    photos = Photo.objects.select_related('point', 'point__road_section', 'inspection').all()

    if point_filter:
        photos = photos.filter(point_id=point_filter)
    if type_filter:
        photos = photos.filter(photo_type=type_filter)
    if road_filter:
        photos = photos.filter(point__road_section_id=road_filter)
    if date_from:
        try:
            d = datetime.strptime(date_from, '%Y-%m-%d').date()
            photos = photos.filter(uploaded_at__date__gte=d)
        except ValueError:
            pass
    if date_to:
        try:
            d = datetime.strptime(date_to, '%Y-%m-%d').date()
            photos = photos.filter(uploaded_at__date__lte=d)
        except ValueError:
            pass

    photos = photos.order_by('-uploaded_at')
    roads = RoadSection.objects.all()
    points = Point.objects.all().order_by('code')

    return render(request, 'roads/photo_list.html', {
        'photos': photos,
        'roads': roads,
        'points': points,
        'point_filter': point_filter,
        'type_filter': type_filter,
        'road_filter': road_filter,
        'date_from': date_from,
        'date_to': date_to,
    })


def photo_upload(request):
    if request.method == 'POST':
        form = PhotoForm(request.POST, request.FILES)
        if form.is_valid():
            photo = form.save()
            if not photo.taken_at:
                photo.taken_at = timezone.now().date()
                photo.save()
            messages.success(request, '照片上传成功。')
            return redirect('roads:photo_list')
    else:
        form = PhotoForm(initial={'taken_at': timezone.now().date()})
    return render(request, 'roads/photo_upload.html', {'form': form})


def photo_detail(request, pk):
    photo = get_object_or_404(
        Photo.objects.select_related('point', 'point__road_section', 'inspection'),
        pk=pk
    )
    related_photos = Photo.objects.filter(point=photo.point).exclude(pk=pk).order_by('-uploaded_at')[:6]
    return render(request, 'roads/photo_detail.html', {
        'photo': photo,
        'related_photos': related_photos,
    })


def photo_delete(request, pk):
    photo = get_object_or_404(Photo, pk=pk)
    if request.method == 'POST':
        photo.delete()
        messages.success(request, '照片已删除。')
        return redirect('roads:photo_list')
    return render(request, 'roads/photo_confirm_delete.html', {'photo': photo})


# ==================== 增强地图：时间轴磨损演变 ====================

def api_points_geo_timeline(request):
    date_str = request.GET.get('date')
    high_risk = request.GET.get('high_risk', '0') == '1'

    points = annotate_high_risk(Point.objects.select_related('road_section').all())
    if high_risk and not date_str:
        points = points.filter(is_high_risk_annotated=True)
    elif high_risk and date_str:
        points = points.filter(
            inspections__wear_level__gte=3,
            inspections__handled=False
        ).distinct()

    target_date = None
    if date_str:
        try:
            target_date = datetime.strptime(date_str, '%Y-%m-%d').date()
        except ValueError:
            target_date = None

    features = []
    for point in points:
        if target_date:
            inspections_before = point.inspections.filter(
                inspection_date__lte=target_date
            ).order_by('-inspection_date')
            latest_at_time = inspections_before.first()
            unhandled_at_time = inspections_before.filter(handled=False).order_by('-wear_level', '-inspection_date').first()
            wear_level = unhandled_at_time.wear_level if unhandled_at_time else None
            is_high = unhandled_at_time and unhandled_at_time.wear_level >= 3
            latest_wear = latest_at_time.wear_level if latest_at_time else None
            latest_wear_display = dict(WEAR_LEVEL_CHOICES).get(latest_wear, '无记录') if latest_at_time else '无记录'
            handled = unhandled_at_time is None
            wear_display = dict(WEAR_LEVEL_CHOICES).get(wear_level, '无未处理记录')
            if high_risk and not is_high:
                continue
        else:
            wear_level = point.worst_unhandled_wear_annotated
            is_high = bool(point.is_high_risk_annotated)
            latest = point.get_latest_inspection()
            latest_wear = latest.wear_level if latest else None
            latest_wear_display = dict(WEAR_LEVEL_CHOICES).get(latest_wear, '无记录') if latest else '无记录'
            handled = wear_level is None
            wear_display = dict(WEAR_LEVEL_CHOICES).get(wear_level, '无未处理记录')

        features.append({
            'type': 'Feature',
            'geometry': {
                'type': 'Point',
                'coordinates': [point.longitude, point.latitude],
            },
            'properties': {
                'id': point.id,
                'code': point.code,
                'name': point.name,
                'point_type': point.point_type,
                'point_type_display': point.get_point_type_display(),
                'road_section': point.road_section.name,
                'road_section_id': point.road_section.id,
                'wear_level': wear_level,
                'wear_level_display': wear_display,
                'is_high_risk': is_high,
                'handled': handled,
                'has_latest': latest_wear is not None,
                'latest_wear_level': latest_wear,
                'latest_wear_display': latest_wear_display,
            }
        })

    return JsonResponse({
        'type': 'FeatureCollection',
        'features': features,
    })


# ==================== 趋势分析：点位磨损演变追踪 ====================

def api_point_wear_history(request, point_pk):
    point = get_object_or_404(Point, pk=point_pk)
    inspections = point.inspections.order_by('inspection_date')

    labels = []
    wear_data = []
    inspections_detail = []
    for insp in inspections:
        labels.append(insp.inspection_date.strftime('%Y-%m-%d'))
        wear_data.append(insp.wear_level)
        inspections_detail.append({
            'id': insp.id,
            'inspection_date': insp.inspection_date.strftime('%Y-%m-%d'),
            'wear_level': insp.wear_level,
            'wear_level_display': dict(WEAR_LEVEL_CHOICES).get(insp.wear_level, '未知'),
            'inspector': insp.inspector,
            'handled': insp.handled,
        })

    return JsonResponse({
        'point': {
            'id': point.id,
            'code': point.code,
            'name': point.name,
            'road': point.road_section.name,
        },
        'labels': labels,
        'wear_data': wear_data,
        'wear_level_choices': dict(WEAR_LEVEL_CHOICES),
        'inspections': inspections_detail,
    })


# ==================== 维护优先级清单 ====================

def priority_list(request):
    unhandled = InspectionRecord.objects.filter(handled=False).select_related(
        'point', 'point__road_section'
    ).order_by('-wear_level', 'inspection_date')

    active_hazards = Hazard.objects.filter(
        status__in=['reported', 'assessing', 'disposing', 'monitoring']
    ).select_related('point', 'point__road_section').order_by('-reported_at')

    priority_items = []

    for insp in unhandled:
        days_pending = (timezone.now().date() - insp.inspection_date).days
        priority_score = insp.wear_level * 100 + days_pending

        if insp.wear_level == 4:
            priority = '紧急'
            priority_class = 'danger'
        elif insp.wear_level == 3:
            priority = '高'
            priority_class = 'warning'
        elif insp.wear_level == 2:
            priority = '中'
            priority_class = 'info'
        else:
            priority = '低'
            priority_class = 'secondary'

        has_task = TaskOrder.objects.filter(
            inspection=insp
        ).exclude(status='closed').first()

        priority_items.append({
            'type': 'wear',
            'type_label': '磨损养护',
            'type_badge': 'bg-primary',
            'inspection': insp,
            'priority': priority,
            'priority_class': priority_class,
            'priority_score': priority_score,
            'days_pending': days_pending,
            'has_task': has_task,
            'point': insp.point,
            'road_section': insp.point.road_section,
            'description': insp.maintenance_suggestion or '定期巡查维护',
            'detail_url': reverse('roads:inspection_detail', args=[insp.pk]),
            'task_create_url': reverse('roads:task_create_from_insp', args=[insp.pk]),
        })

    hazard_level_score = {
        'critical': 500,
        'severe': 400,
        'warning': 300,
        'info': 200,
        'safe': 100,
    }
    hazard_priority_map = {
        'critical': ('紧急', 'danger'),
        'severe': ('高', 'warning'),
        'warning': ('中', 'info'),
        'info': ('低', 'secondary'),
        'safe': ('低', 'secondary'),
    }

    for hazard in active_hazards:
        days_pending = (timezone.now().date() - hazard.reported_at.date()).days
        level_score = hazard_level_score.get(hazard.hazard_level, 100)
        priority_score = level_score + days_pending

        priority, priority_class = hazard_priority_map.get(
            hazard.hazard_level, ('低', 'secondary')
        )

        has_task = TaskOrder.objects.filter(
            hazard=hazard
        ).exclude(status='closed').first()

        priority_items.append({
            'type': 'hazard',
            'type_label': '灾害隐患',
            'type_badge': 'bg-danger',
            'hazard': hazard,
            'priority': priority,
            'priority_class': priority_class,
            'priority_score': priority_score,
            'days_pending': days_pending,
            'has_task': has_task,
            'point': hazard.point,
            'road_section': hazard.point.road_section,
            'description': hazard.description or hazard.get_hazard_type_display(),
            'detail_url': reverse('roads:hazard_detail', args=[hazard.pk]),
            'task_create_url': reverse('roads:hazard_task_create', args=[hazard.pk]),
        })

    priority_items.sort(key=lambda x: x['priority_score'], reverse=True)

    stats = {
        'urgent': sum(1 for p in priority_items if p['priority'] == '紧急'),
        'high': sum(1 for p in priority_items if p['priority'] == '高'),
        'medium': sum(1 for p in priority_items if p['priority'] == '中'),
        'low': sum(1 for p in priority_items if p['priority'] == '低'),
        'with_task': sum(1 for p in priority_items if p['has_task']),
        'total': len(priority_items),
        'wear_count': sum(1 for p in priority_items if p['type'] == 'wear'),
        'hazard_count': sum(1 for p in priority_items if p['type'] == 'hazard'),
    }
    stats['without_task'] = stats['total'] - stats['with_task']

    return render(request, 'roads/priority_list.html', {
        'priority_items': priority_items,
        'stats': stats,
    })


# ==================== 数据导出模块 ====================

def export_page(request):
    if request.method == 'POST':
        form = DataExportForm(request.POST)
        if form.is_valid():
            return _do_export(form)
    else:
        form = DataExportForm()
    return render(request, 'roads/export.html', {'form': form})


def _do_export(form):
    export_type = form.cleaned_data['export_type']
    export_format = form.cleaned_data['export_format']
    road_section = form.cleaned_data.get('road_section')
    date_from = form.cleaned_data.get('date_from')
    date_to = form.cleaned_data.get('date_to')

    if export_type == 'inspections':
        qs = InspectionRecord.objects.select_related('point', 'point__road_section').all()
        if road_section:
            qs = qs.filter(point__road_section=road_section)
        if date_from:
            qs = qs.filter(inspection_date__gte=date_from)
        if date_to:
            qs = qs.filter(inspection_date__lte=date_to)
        headers = ['路段', '点位编号', '点位名称', '巡查日期', '巡查人员', '磨损等级', '磨损描述', '维护建议', '是否已处理', '处理日期']
        rows = []
        for r in qs:
            rows.append([
                r.point.road_section.name,
                r.point.code,
                r.point.name,
                r.inspection_date.strftime('%Y-%m-%d'),
                r.inspector,
                r.get_wear_level_display(),
                r.wear_description or '',
                r.maintenance_suggestion or '',
                '是' if r.handled else '否',
                r.handled_date.strftime('%Y-%m-%d') if r.handled_date else '',
            ])
        filename = f'inspections_{timezone.now().strftime("%Y%m%d%H%M%S")}'

    elif export_type == 'points':
        qs = Point.objects.select_related('road_section').all()
        if road_section:
            qs = qs.filter(road_section=road_section)
        headers = ['路段', '点位编号', '点位名称', '类型', '纬度', '经度', '描述', '当前磨损等级', '是否高风险']
        rows = []
        for p in qs:
            wl = p.get_current_wear_level()
            rows.append([
                p.road_section.name,
                p.code,
                p.name,
                p.get_point_type_display(),
                p.latitude,
                p.longitude,
                p.description or '',
                dict(WEAR_LEVEL_CHOICES).get(wl, '无记录'),
                '是' if p.is_high_risk() else '否',
            ])
        filename = f'points_{timezone.now().strftime("%Y%m%d%H%M%S")}'

    elif export_type == 'tasks':
        qs = TaskOrder.objects.select_related('point', 'point__road_section', 'inspection').all()
        if road_section:
            qs = qs.filter(point__road_section=road_section)
        if date_from:
            qs = qs.filter(created_at__date__gte=date_from)
        if date_to:
            qs = qs.filter(created_at__date__lte=date_to)
        headers = ['工单标题', '路段', '点位', '状态', '优先级', '指派人员', '整改期限', '创建时间', '派发时间', '整改完成时间', '复核时间', '闭环时间', '是否超期']
        rows = []
        for t in qs:
            rows.append([
                t.title,
                t.point.road_section.name,
                t.point.code + ' ' + t.point.name,
                t.get_status_display(),
                t.get_priority_display(),
                t.assigned_to or '',
                t.deadline.strftime('%Y-%m-%d') if t.deadline else '',
                t.created_at.strftime('%Y-%m-%d %H:%M'),
                t.dispatched_at.strftime('%Y-%m-%d %H:%M') if t.dispatched_at else '',
                t.rectified_at.strftime('%Y-%m-%d %H:%M') if t.rectified_at else '',
                t.reviewed_at.strftime('%Y-%m-%d %H:%M') if t.reviewed_at else '',
                t.closed_at.strftime('%Y-%m-%d %H:%M') if t.closed_at else '',
                '是' if t.is_overdue() else '否',
            ])
        filename = f'tasks_{timezone.now().strftime("%Y%m%d%H%M%S")}'

    else:
        qs = Alert.objects.select_related('point', 'point__road_section').all()
        if road_section:
            qs = qs.filter(point__road_section=road_section)
        if date_from:
            qs = qs.filter(created_at__date__gte=date_from)
        if date_to:
            qs = qs.filter(created_at__date__lte=date_to)
        headers = ['预警类型', '预警等级', '路段', '点位', '预警信息', '是否已读', '是否已处理', '创建时间', '处理时间']
        rows = []
        for a in qs:
            rows.append([
                a.get_alert_type_display(),
                a.get_alert_level_display(),
                a.point.road_section.name,
                a.point.code + ' ' + a.point.name,
                a.message,
                '是' if a.is_read else '否',
                '是' if a.is_resolved else '否',
                a.created_at.strftime('%Y-%m-%d %H:%M'),
                a.resolved_at.strftime('%Y-%m-%d %H:%M') if a.resolved_at else '',
            ])
        filename = f'alerts_{timezone.now().strftime("%Y%m%d%H%M%S")}'

    if export_format == 'csv':
        return _export_csv(headers, rows, filename)
    else:
        return _export_xlsx(headers, rows, filename)


def _export_csv(headers, rows, filename):
    buffer = io.StringIO()
    buffer.write('\ufeff')
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows(rows)
    response = HttpResponse(buffer.getvalue(), content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
    return response


def _export_xlsx(headers, rows, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = '数据导出'

    header_fill = PatternFill(start_color='5c4033', end_color='5c4033', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical='center')
            cell.border = thin_border

    for col_idx in range(1, len(headers) + 1):
        max_length = len(str(headers[col_idx - 1]))
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
            for cell in row:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_length + 4, 50)

    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    return response


# ==================== 灾害隐患与通行安全预警模块 ====================

def _sync_road_passage_status(road_section):
    if not road_section:
        return
    ps, created = RoadPassageStatus.objects.get_or_create(
        road_section=road_section
    )
    ps.recalculate_from_hazards()


def hazard_list(request):
    level_filter = request.GET.get('level')
    type_filter = request.GET.get('type')
    loc_filter = request.GET.get('loc_type')
    status_filter = request.GET.get('status')
    passage_filter = request.GET.get('passage')
    road_filter = request.GET.get('road')
    overdue_filter = request.GET.get('overdue')
    keyword = request.GET.get('q')

    hazards = Hazard.objects.select_related(
        'road_section', 'point', 'inspection_source'
    ).prefetch_related('disposals').all()

    if level_filter:
        hazards = hazards.filter(hazard_level=level_filter)
    if type_filter:
        hazards = hazards.filter(hazard_type=type_filter)
    if loc_filter:
        hazards = hazards.filter(location_type=loc_filter)
    if status_filter:
        hazards = hazards.filter(status=status_filter)
    if passage_filter:
        hazards = hazards.filter(passage_status=passage_filter)
    if road_filter:
        hazards = hazards.filter(road_section_id=road_filter)
    if overdue_filter == '1':
        today = timezone.now().date()
        hazards = hazards.exclude(
            status__in=['resolved', 'closed']
        ).filter(disposal_deadline__lt=today)
    if keyword:
        hazards = hazards.filter(
            Q(title__icontains=keyword) |
            Q(code__icontains=keyword) |
            Q(description__icontains=keyword) |
            Q(location_desc__icontains=keyword)
        )

    hazards = hazards.order_by('-reported_at')
    roads = RoadSection.objects.all()

    stats = {
        'total': hazards.count(),
        'critical': hazards.filter(hazard_level='critical').count(),
        'severe': hazards.filter(hazard_level='severe').count(),
        'warning': hazards.filter(hazard_level='warning').count(),
        'reported': hazards.filter(status='reported').count(),
        'disposing': hazards.filter(status='disposing').count(),
        'resolved': hazards.filter(status__in=['resolved', 'closed']).count(),
        'overdue': hazards.exclude(
            status__in=['resolved', 'closed']
        ).filter(disposal_deadline__lt=timezone.now().date()).count(),
    }

    return render(request, 'roads/hazard_list.html', {
        'hazards': hazards,
        'roads': roads,
        'level_filter': level_filter,
        'type_filter': type_filter,
        'loc_filter': loc_filter,
        'status_filter': status_filter,
        'passage_filter': passage_filter,
        'road_filter': road_filter,
        'overdue_filter': overdue_filter,
        'keyword': keyword or '',
        'hazard_level_choices': dict(HAZARD_LEVEL_CHOICES),
        'hazard_type_choices': dict(HAZARD_TYPE_CHOICES),
        'location_type_choices': dict(HAZARD_LOCATION_TYPE_CHOICES),
        'hazard_status_choices': dict(HAZARD_STATUS_CHOICES),
        'passage_status_choices': dict(PASSAGE_STATUS_CHOICES),
        'stats': stats,
    })


def hazard_create(request):
    if request.method == 'POST':
        form = HazardForm(request.POST)
        if form.is_valid():
            try:
                hazard = form.save(commit=False)
                hazard.status = 'reported'
                hazard.save()
                _sync_road_passage_status(hazard.road_section)
                messages.success(request, f'隐患 {hazard.code} 上报成功。')
                return redirect('roads:hazard_detail', pk=hazard.pk)
            except ValidationError as e:
                for field, errors in e.message_dict.items():
                    for error in errors:
                        form.add_error(field, error)
    else:
        form = HazardForm(initial={
            'reported_date': timezone.now().date(),
            'reported_by': request.user.username if hasattr(request, 'user') and request.user.is_authenticated else '系统',
        })
    return render(request, 'roads/hazard_form.html', {'form': form, 'mode': 'create'})


def hazard_detail(request, pk):
    hazard = get_object_or_404(
        Hazard.objects.select_related(
            'road_section', 'point', 'inspection_source'
        ).prefetch_related('disposals', 'disposals__related_task', 'task_orders'),
        pk=pk
    )
    disposals = hazard.disposals.all().order_by('-disposed_at')
    task_orders = hazard.task_orders.all().order_by('-created_at')
    return render(request, 'roads/hazard_detail.html', {
        'hazard': hazard,
        'disposals': disposals,
        'task_orders': task_orders,
        'hazard_level_choices': dict(HAZARD_LEVEL_CHOICES),
        'hazard_type_choices': dict(HAZARD_TYPE_CHOICES),
        'location_type_choices': dict(HAZARD_LOCATION_TYPE_CHOICES),
        'hazard_status_choices': dict(HAZARD_STATUS_CHOICES),
        'passage_status_choices': dict(PASSAGE_STATUS_CHOICES),
        'control_choices': dict(CONTROL_SUGGESTION_CHOICES),
        'disposal_type_choices': dict(DISPOSAL_TYPE_CHOICES),
        'task_status_choices': dict(TASK_STATUS_CHOICES),
    })


def hazard_edit(request, pk):
    hazard = get_object_or_404(Hazard, pk=pk)
    if request.method == 'POST':
        form = HazardForm(request.POST, instance=hazard)
        if form.is_valid():
            try:
                h = form.save()
                _sync_road_passage_status(h.road_section)
                messages.success(request, f'隐患 {hazard.code} 更新成功。')
                return redirect('roads:hazard_detail', pk=hazard.pk)
            except ValidationError as e:
                for field, errors in e.message_dict.items():
                    for error in errors:
                        form.add_error(field, error)
    else:
        form = HazardForm(instance=hazard)
    return render(request, 'roads/hazard_form.html', {
        'form': form, 'hazard': hazard, 'mode': 'edit'
    })


def hazard_assess(request, pk):
    hazard = get_object_or_404(Hazard, pk=pk)
    if hazard.status not in ['reported', 'assessing']:
        messages.error(request, '当前状态不允许评估。')
        return redirect('roads:hazard_detail', pk=pk)

    if request.method == 'POST':
        form = HazardAssessForm(request.POST, instance=hazard)
        if form.is_valid():
            h = form.save(commit=False)
            if not h.assessed_at:
                h.assessed_at = timezone.now()
            if h.status == 'reported':
                h.status = 'pending_disposal'
            h.save()
            HazardDisposal.objects.create(
                hazard=h,
                disposal_type='assessed',
                status_before='reported',
                status_after=h.status,
                level_before='info',
                level_after=h.hazard_level,
                passage_before='normal',
                passage_after=h.passage_status,
                description=h.assess_note or '完成隐患等级评估',
                disposed_by=h.assessed_by or '系统',
                disposal_result=f'评估等级：{h.get_hazard_level_display()}，封控建议：{h.get_control_suggestion_display()}',
                next_step='按处置方案开展处置工作'
            )
            _sync_road_passage_status(h.road_section)
            messages.success(request, '隐患评估完成。')
            return redirect('roads:hazard_detail', pk=pk)
    else:
        form = HazardAssessForm(instance=hazard, initial={
            'status': 'pending_disposal' if hazard.status == 'reported' else hazard.status,
            'disposal_deadline': hazard.disposal_deadline or (timezone.now().date() + timedelta(days=7)),
        })
    return render(request, 'roads/hazard_assess.html', {
        'form': form, 'hazard': hazard
    })


def hazard_update_status(request, pk):
    hazard = get_object_or_404(Hazard, pk=pk)
    if request.method == 'POST':
        form = HazardStatusForm(request.POST, instance=hazard)
        if form.is_valid():
            old_status = hazard.status
            old_level = hazard.hazard_level
            old_passage = hazard.passage_status
            h = form.save()
            HazardDisposal.objects.create(
                hazard=h,
                disposal_type='other',
                status_before=old_status,
                status_after=h.status,
                level_before=old_level,
                level_after=h.hazard_level,
                passage_before=old_passage,
                passage_after=h.passage_status,
                description=request.POST.get('note', '更新隐患状态'),
                disposed_by=request.POST.get('operator', '系统'),
            )
            _sync_road_passage_status(h.road_section)
            messages.success(request, '隐患状态已更新。')
        else:
            messages.error(request, '表单验证失败。')
    return redirect(request.META.get('HTTP_REFERER', 'roads:hazard_detail'), pk=pk)


def hazard_dispose(request, pk):
    hazard = get_object_or_404(Hazard, pk=pk)
    if hazard.status in ['resolved', 'closed']:
        messages.warning(request, '该隐患已闭环，无需再次处置。')
        return redirect('roads:hazard_detail', pk=pk)

    if request.method == 'POST':
        form = HazardDisposalForm(request.POST)
        if form.is_valid():
            d = form.save(commit=False)
            d.hazard = hazard
            d.status_before = hazard.status
            d.level_before = hazard.hazard_level
            d.passage_before = hazard.passage_status
            d.save()

            if d.status_after:
                hazard.status = d.status_after
            if d.level_after:
                hazard.hazard_level = d.level_after
            if d.passage_after:
                hazard.passage_status = d.passage_after
            hazard.save()
            _sync_road_passage_status(hazard.road_section)

            messages.success(request, '处置记录已添加。')
            return redirect('roads:hazard_detail', pk=pk)
    else:
        form = HazardDisposalForm(initial={
            'disposed_at': timezone.now(),
            'disposal_type': 'onsite',
            'status_after': 'disposing' if hazard.status not in ['disposing', 'monitoring'] else hazard.status,
        })
    return render(request, 'roads/hazard_dispose.html', {
        'form': form, 'hazard': hazard
    })


def hazard_close(request, pk):
    hazard = get_object_or_404(Hazard, pk=pk)
    if request.method == 'POST':
        old_status = hazard.status
        closed_by = request.POST.get('closed_by', '系统')
        note = request.POST.get('close_note', '')
        hazard.status = 'closed'
        hazard.closed_by = closed_by
        hazard.save()
        HazardDisposal.objects.create(
            hazard=hazard,
            disposal_type='closed',
            status_before=old_status,
            status_after='closed',
            level_before=hazard.hazard_level,
            level_after='safe',
            passage_before=hazard.passage_status,
            passage_after='normal',
            description=note or '隐患已消除，完成闭环',
            disposed_by=closed_by,
            disposal_result='隐患闭环归档'
        )
        _sync_road_passage_status(hazard.road_section)
        messages.success(request, '隐患已闭环归档。')
    return redirect('roads:hazard_detail', pk=pk)


def hazard_delete(request, pk):
    hazard = get_object_or_404(Hazard, pk=pk)
    road_section = hazard.road_section
    if request.method == 'POST':
        hazard.delete()
        _sync_road_passage_status(road_section)
        messages.success(request, '隐患记录已删除。')
        return redirect('roads:hazard_list')
    return render(request, 'roads/hazard_confirm_delete.html', {'hazard': hazard})


def hazard_task_create(request, pk):
    hazard = get_object_or_404(Hazard, pk=pk)

    level_priority_map = {
        'critical': 'urgent',
        'severe': 'urgent',
        'warning': 'high',
        'info': 'medium',
        'safe': 'low',
    }
    priority = level_priority_map.get(hazard.hazard_level, 'high')

    initial = {
        'hazard': hazard,
        'title': f'{hazard.code} - {hazard.title} 应急处置',
        'description': f'隐患编号：{hazard.code}\n隐患类型：{hazard.get_hazard_type_display()}\n预警等级：{hazard.get_hazard_level_display()}\n\n隐患描述：\n{hazard.description}\n\n封控建议：{hazard.get_control_suggestion_display()}',
        'priority': priority,
    }
    if hazard.point:
        initial['point'] = hazard.point
    if hazard.disposal_deadline:
        initial['deadline'] = hazard.disposal_deadline
    else:
        days_map = {'critical': 1, 'severe': 3, 'warning': 7, 'info': 14, 'safe': 30}
        initial['deadline'] = timezone.now().date() + timedelta(days=days_map.get(hazard.hazard_level, 7))

    if request.method == 'POST':
        form = TaskOrderForm(request.POST)
        if form.is_valid():
            task = form.save(commit=False)
            task.hazard = hazard
            task.status = 'dispatched'
            task.dispatched_at = timezone.now()
            task.save()

            disposal = HazardDisposal.objects.create(
                hazard=hazard,
                disposal_type='dispatch',
                status_before=hazard.status,
                status_after='disposing' if hazard.status in ['reported', 'assessing'] else hazard.status,
                level_before=hazard.hazard_level,
                passage_before=hazard.passage_status,
                description=f'派发应急处置工单：{task.title}\n指派人员：{task.assigned_to or "待指派"}\n整改期限：{task.deadline or "未设定"}',
                disposed_by=request.user.username if request.user.is_authenticated else '系统',
                related_task=task,
            )

            if hazard.status in ['reported', 'assessing']:
                hazard.status = 'disposing'
                hazard.save()

            messages.success(request, f'应急工单已创建并派发：{task.title}')
            return redirect('roads:hazard_detail', pk=pk)
    else:
        form = TaskOrderForm(initial=initial)
    return render(request, 'roads/hazard_task_create.html', {
        'form': form,
        'hazard': hazard,
    })


# ==================== 通行状态管理 ====================

def passage_status_list(request):
    road_filter = request.GET.get('road')

    for road in RoadSection.objects.all():
        RoadPassageStatus.objects.get_or_create(road_section=road)

    qs = RoadPassageStatus.objects.select_related(
        'road_section'
    ).all().order_by('road_section__code')

    if road_filter:
        qs = qs.filter(road_section_id=road_filter)

    roads = RoadSection.objects.all()

    stats = {
        'total': qs.count(),
        'normal': qs.filter(passage_status='normal').count(),
        'caution': qs.filter(passage_status='caution').count(),
        'restricted': qs.filter(passage_status='restricted').count(),
        'detour': qs.filter(passage_status='detour').count(),
        'closed': qs.filter(passage_status='closed').count(),
        'with_hazards': qs.filter(active_hazard_count__gt=0).count(),
    }

    return render(request, 'roads/passage_status_list.html', {
        'passage_statuses': qs,
        'roads': roads,
        'road_filter': road_filter,
        'passage_status_choices': dict(PASSAGE_STATUS_CHOICES),
        'stats': stats,
    })


def passage_status_edit(request, pk):
    ps = get_object_or_404(
        RoadPassageStatus.objects.select_related('road_section'),
        pk=pk
    )
    if request.method == 'POST':
        form = RoadPassageStatusForm(request.POST, instance=ps)
        if form.is_valid():
            form.save()
            messages.success(request, '路段通行状态已更新。')
            return redirect('roads:passage_status_list')
    else:
        form = RoadPassageStatusForm(instance=ps)
    return render(request, 'roads/passage_status_form.html', {
        'form': form, 'passage_status': ps
    })


def passage_status_recalc(request, pk):
    ps = get_object_or_404(RoadPassageStatus, pk=pk)
    old_status = ps.passage_status
    new_status = ps.recalculate_from_hazards()
    if old_status != new_status:
        messages.info(request, f'通行状态已自动更新：从{ps.get_passage_status_display()}调整为{dict(PASSAGE_STATUS_CHOICES)[new_status]}')
    else:
        messages.success(request, '已重新计算，状态无变化。')
    return redirect(request.META.get('HTTP_REFERER', 'roads:passage_status_list'))


# ==================== 灾害隐患地图与API ====================

def hazard_map(request):
    level_filter = request.GET.get('level')
    status_filter = request.GET.get('status')
    road_filter = request.GET.get('road')

    return render(request, 'roads/hazard_map.html', {
        'level_filter': level_filter,
        'status_filter': status_filter,
        'road_filter': road_filter,
        'hazard_level_choices': dict(HAZARD_LEVEL_CHOICES),
        'hazard_status_choices': dict(HAZARD_STATUS_CHOICES),
        'passage_status_choices': dict(PASSAGE_STATUS_CHOICES),
        'roads': RoadSection.objects.all(),
    })


def api_hazards_geo(request):
    level_filter = request.GET.get('level')
    status_filter = request.GET.get('status')
    road_filter = request.GET.get('road')

    hazards = Hazard.objects.select_related(
        'road_section', 'point'
    ).exclude(
        status__in=['resolved', 'closed']
    ).filter(
        latitude__isnull=False,
        longitude__isnull=False
    )

    if level_filter:
        hazards = hazards.filter(hazard_level=level_filter)
    if status_filter:
        hazards = hazards.filter(status=status_filter)
    if road_filter:
        hazards = hazards.filter(road_section_id=road_filter)

    features = []
    for hz in hazards:
        features.append({
            'type': 'Feature',
            'geometry': {
                'type': 'Point',
                'coordinates': [hz.longitude, hz.latitude],
            },
            'properties': {
                'id': hz.id,
                'code': hz.code,
                'title': hz.title,
                'location_type': hz.location_type,
                'location_type_display': hz.get_location_type_display(),
                'hazard_type': hz.hazard_type,
                'hazard_type_display': hz.get_hazard_type_display(),
                'hazard_level': hz.hazard_level,
                'hazard_level_display': hz.get_hazard_level_display(),
                'hazard_level_color': hz.get_hazard_level_color(),
                'status': hz.status,
                'status_display': hz.get_status_display(),
                'passage_status': hz.passage_status,
                'passage_display': hz.get_passage_status_display(),
                'passage_color': hz.get_passage_status_color(),
                'road_section': hz.road_section.name if hz.road_section else '',
                'road_section_id': hz.road_section.id if hz.road_section else None,
                'reported_at': hz.reported_at.strftime('%Y-%m-%d %H:%M'),
                'reported_by': hz.reported_by,
                'control_suggestion': hz.get_control_suggestion_display(),
                'progress': hz.get_progress_percentage(),
                'description': hz.description[:100] + ('...' if len(hz.description) > 100 else ''),
                'is_high_risk': hz.hazard_level in ['severe', 'critical'],
            }
        })

    return JsonResponse({
        'type': 'FeatureCollection',
        'features': features,
    })


def api_roads_passage_status(request):
    result = []
    for road in RoadSection.objects.all():
        ps, _ = RoadPassageStatus.objects.get_or_create(road_section=road)
        hazards = road.hazards.exclude(status__in=['resolved', 'closed'])
        hazard_points = []
        for hz in hazards.filter(latitude__isnull=False, longitude__isnull=False):
            hazard_points.append([hz.longitude, hz.latitude])
        result.append({
            'road_id': road.id,
            'road_code': road.code,
            'road_name': road.name,
            'passage_status': ps.passage_status,
            'passage_display': ps.get_passage_status_display(),
            'status_color': ps.get_status_color(),
            'active_hazard_count': ps.active_hazard_count,
            'critical_hazard_count': ps.critical_hazard_count,
            'status_reason': ps.status_reason,
            'hazard_points': hazard_points,
        })
    return JsonResponse({'roads': result})


def api_hazard_summary(request):
    total = Hazard.objects.count()
    active = Hazard.objects.exclude(status__in=['resolved', 'closed'])
    today = timezone.now().date()
    level_data = {}
    for level, label in HAZARD_LEVEL_CHOICES:
        level_data[level] = {
            'label': label,
            'active': active.filter(hazard_level=level).count(),
            'total': Hazard.objects.filter(hazard_level=level).count(),
        }
    status_data = {}
    for status, label in HAZARD_STATUS_CHOICES:
        status_data[status] = {
            'label': label,
            'count': Hazard.objects.filter(status=status).count(),
        }
    type_data = {}
    for htype, label in HAZARD_TYPE_CHOICES:
        type_data[htype] = {
            'label': label,
            'active': active.filter(hazard_type=htype).count(),
        }
    overdue = active.filter(disposal_deadline__lt=today).count()
    today_new = Hazard.objects.filter(reported_date=today).count()

    return JsonResponse({
        'total': total,
        'active': active.count(),
        'resolved': total - active.count(),
        'overdue': overdue,
        'today_new': today_new,
        'by_level': level_data,
        'by_status': status_data,
        'by_type': type_data,
    })


# ==================== 游客承载与开放时段管理模块 ====================

def _get_current_season():
    month = timezone.now().month
    if month in [3, 4, 5]:
        return 'spring'
    elif month in [6, 7, 8]:
        return 'summer'
    elif month in [9, 10, 11]:
        return 'autumn'
    return 'winter'


def _generate_open_suggestion(road_section):
    suggestions = []
    schedule = OpenSchedule.objects.filter(
        road_section=road_section, is_active=True
    ).first()
    passage = RoadPassageStatus.objects.filter(
        road_section=road_section
    ).first()
    active_controls = TemporaryControl.objects.filter(
        road_section=road_section, is_active=True,
        end_time__gte=timezone.now()
    )
    hazards = road_section.hazards.exclude(status__in=['resolved', 'closed'])
    critical_hazards = hazards.filter(hazard_level__in=['severe', 'critical'])

    if critical_hazards.exists():
        suggestions.append({
            'level': 'critical',
            'message': f'存在{critical_hazards.count()}处紧急/严重隐患，建议暂停开放。',
            'action': 'closed'
        })
    elif hazards.filter(hazard_level='warning').exists():
        suggestions.append({
            'level': 'warning',
            'message': '存在警告级隐患，建议限制开放并降低承载量。',
            'action': 'restricted'
        })

    if passage and passage.passage_status == 'closed':
        suggestions.append({
            'level': 'critical',
            'message': '路段通行状态为禁止通行，不宜开放游览。',
            'action': 'closed'
        })
    elif passage and passage.passage_status in ['restricted', 'detour']:
        suggestions.append({
            'level': 'warning',
            'message': f'路段通行状态为{passage.get_passage_status_display()}，建议限制开放。',
            'action': 'restricted'
        })

    if schedule and schedule.maintenance_status != 'normal':
        status_map = {
            'maintaining': '养护中',
            'repairing': '修缮中',
            'emergency': '应急抢修'
        }
        suggestions.append({
            'level': 'warning',
            'message': f'路段处于{status_map.get(schedule.maintenance_status, "异常")}状态，建议调整开放计划。',
            'action': 'restricted' if schedule.maintenance_status != 'emergency' else 'closed'
        })

    unhandled_critical = road_section.points.filter(
        inspections__wear_level=4, inspections__handled=False
    ).exists()
    if unhandled_critical:
        suggestions.append({
            'level': 'warning',
            'message': '存在未处理的严重磨损点位，建议控制游客量。',
            'action': 'restricted'
        })

    if active_controls.exists():
        for ctrl in active_controls:
            suggestions.append({
                'level': 'info',
                'message': f'临时管制：{ctrl.get_control_type_display()} - {ctrl.reason[:50]}',
                'action': ctrl.open_status
            })

    if not suggestions:
        suggestions.append({
            'level': 'info',
            'message': '路段状态正常，按计划开放。',
            'action': 'open'
        })

    return suggestions


def open_schedule_list(request):
    road_filter = request.GET.get('road')
    season_filter = request.GET.get('season')
    weather_filter = request.GET.get('weather')
    maintenance_filter = request.GET.get('maintenance')

    schedules = OpenSchedule.objects.select_related('road_section').all()

    if road_filter:
        schedules = schedules.filter(road_section_id=road_filter)
    if season_filter:
        schedules = schedules.filter(season=season_filter)
    if weather_filter:
        schedules = schedules.filter(weather_condition=weather_filter)
    if maintenance_filter:
        schedules = schedules.filter(maintenance_status=maintenance_filter)

    schedules = schedules.order_by('road_section__code', 'season')
    roads = RoadSection.objects.all()

    stats = {
        'total': schedules.count(),
        'active': schedules.filter(is_active=True).count(),
        'inactive': schedules.filter(is_active=False).count(),
    }

    return render(request, 'roads/open_schedule_list.html', {
        'schedules': schedules,
        'roads': roads,
        'road_filter': road_filter,
        'season_filter': season_filter,
        'weather_filter': weather_filter,
        'maintenance_filter': maintenance_filter,
        'season_choices': dict(SEASON_CHOICES),
        'weather_choices': dict(WEATHER_CONDITION_CHOICES),
        'maintenance_choices': dict(MAINTENANCE_STATUS_CHOICES),
        'stats': stats,
    })


def open_schedule_create(request):
    initial = {
        'season': _get_current_season(),
        'time_slot_minutes': 60,
    }
    road_id = request.GET.get('road')
    if road_id:
        initial['road_section'] = road_id

    if request.method == 'POST':
        form = OpenScheduleForm(request.POST)
        if form.is_valid():
            try:
                schedule = form.save(commit=False)
                schedule.full_clean()
                schedule.save()
                messages.success(request, '开放计划创建成功。')
                return redirect('roads:open_schedule_list')
            except ValidationError as e:
                for field, errors in e.message_dict.items():
                    for error in errors:
                        form.add_error(field, error)
    else:
        form = OpenScheduleForm(initial=initial)
    return render(request, 'roads/open_schedule_form.html', {
        'form': form, 'mode': 'create'
    })


def open_schedule_edit(request, pk):
    schedule = get_object_or_404(OpenSchedule, pk=pk)
    if request.method == 'POST':
        form = OpenScheduleForm(request.POST, instance=schedule)
        if form.is_valid():
            try:
                s = form.save(commit=False)
                s.full_clean()
                s.save()
                messages.success(request, '开放计划更新成功。')
                return redirect('roads:open_schedule_list')
            except ValidationError as e:
                for field, errors in e.message_dict.items():
                    for error in errors:
                        form.add_error(field, error)
    else:
        form = OpenScheduleForm(instance=schedule)
    return render(request, 'roads/open_schedule_form.html', {
        'form': form, 'schedule': schedule, 'mode': 'edit'
    })


def open_schedule_delete(request, pk):
    schedule = get_object_or_404(OpenSchedule, pk=pk)
    if request.method == 'POST':
        schedule.delete()
        messages.success(request, '开放计划已删除。')
        return redirect('roads:open_schedule_list')
    return render(request, 'roads/open_schedule_confirm_delete.html', {
        'schedule': schedule
    })


def temporary_control_list(request):
    road_filter = request.GET.get('road')
    type_filter = request.GET.get('type')
    status_filter = request.GET.get('status')

    controls = TemporaryControl.objects.select_related('road_section').all()

    if road_filter:
        controls = controls.filter(road_section_id=road_filter)
    if type_filter:
        controls = controls.filter(control_type=type_filter)
    if status_filter == 'active':
        controls = controls.filter(is_active=True)
    elif status_filter == 'inactive':
        controls = controls.filter(is_active=False)

    controls = controls.order_by('-start_time')
    roads = RoadSection.objects.all()

    stats = {
        'total': controls.count(),
        'active': controls.filter(is_active=True).count(),
        'inactive': controls.filter(is_active=False).count(),
    }

    return render(request, 'roads/temporary_control_list.html', {
        'controls': controls,
        'roads': roads,
        'road_filter': road_filter,
        'type_filter': type_filter,
        'status_filter': status_filter,
        'control_type_choices': dict(CONTROL_TYPE_CHOICES),
        'open_status_choices': dict(OPEN_STATUS_CHOICES),
        'stats': stats,
    })


def temporary_control_create(request):
    initial = {
        'start_time': timezone.now(),
        'end_time': timezone.now() + timedelta(hours=24),
        'issued_by': request.user.username if hasattr(request, 'user') and request.user.is_authenticated else '系统',
    }
    road_id = request.GET.get('road')
    if road_id:
        initial['road_section'] = road_id

    if request.method == 'POST':
        form = TemporaryControlForm(request.POST)
        if form.is_valid():
            try:
                control = form.save(commit=False)
                control.full_clean()
                control.save()
                messages.success(request, '临时管制登记成功。')
                return redirect('roads:temporary_control_list')
            except ValidationError as e:
                for field, errors in e.message_dict.items():
                    for error in errors:
                        form.add_error(field, error)
    else:
        form = TemporaryControlForm(initial=initial)
    return render(request, 'roads/temporary_control_form.html', {
        'form': form, 'mode': 'create'
    })


def temporary_control_edit(request, pk):
    control = get_object_or_404(TemporaryControl, pk=pk)
    if request.method == 'POST':
        form = TemporaryControlForm(request.POST, instance=control)
        if form.is_valid():
            try:
                c = form.save(commit=False)
                c.full_clean()
                c.save()
                messages.success(request, '临时管制更新成功。')
                return redirect('roads:temporary_control_list')
            except ValidationError as e:
                for field, errors in e.message_dict.items():
                    for error in errors:
                        form.add_error(field, error)
    else:
        form = TemporaryControlForm(instance=control)
    return render(request, 'roads/temporary_control_form.html', {
        'form': form, 'control': control, 'mode': 'edit'
    })


def temporary_control_deactivate(request, pk):
    control = get_object_or_404(TemporaryControl, pk=pk)
    if request.method == 'POST':
        control.deactivate()
        messages.success(request, '临时管制已解除。')
    return redirect('roads:temporary_control_list')


def temporary_control_delete(request, pk):
    control = get_object_or_404(TemporaryControl, pk=pk)
    if request.method == 'POST':
        control.delete()
        messages.success(request, '临时管制记录已删除。')
        return redirect('roads:temporary_control_list')
    return render(request, 'roads/temporary_control_confirm_delete.html', {
        'control': control
    })


def visitor_flow_list(request):
    road_filter = request.GET.get('road')
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')
    type_filter = request.GET.get('type')
    weather_filter = request.GET.get('weather')

    flows = VisitorFlowRecord.objects.select_related('road_section').all()

    if road_filter:
        flows = flows.filter(road_section_id=road_filter)
    if date_from:
        try:
            d = datetime.strptime(date_from, '%Y-%m-%d').date()
            flows = flows.filter(record_date__gte=d)
        except ValueError:
            pass
    if date_to:
        try:
            d = datetime.strptime(date_to, '%Y-%m-%d').date()
            flows = flows.filter(record_date__lte=d)
        except ValueError:
            pass
    if type_filter:
        flows = flows.filter(record_type=type_filter)
    if weather_filter:
        flows = flows.filter(weather=weather_filter)

    flows = flows.order_by('-record_date', '-record_time')
    roads = RoadSection.objects.all()

    today = timezone.now().date()
    today_flows = VisitorFlowRecord.objects.filter(record_date=today)
    today_total = today_flows.aggregate(total=Count('id'))['total']
    today_visitors = today_flows.aggregate(
        total=Count('visitor_count')
    )['total'] if today_flows.exists() else 0

    latest_occupancy = {}
    for road in roads:
        latest = VisitorFlowRecord.objects.filter(
            road_section=road
        ).order_by('-record_date', '-record_time').first()
        if latest:
            latest_occupancy[road.pk] = latest.current_occupancy

    stats = {
        'total_records': flows.count(),
        'today_records': today_total,
        'today_visitors': today_visitors,
    }

    return render(request, 'roads/visitor_flow_list.html', {
        'flows': flows,
        'roads': roads,
        'road_filter': road_filter,
        'date_from': date_from,
        'date_to': date_to,
        'type_filter': type_filter,
        'weather_filter': weather_filter,
        'flow_type_choices': dict(FLOW_RECORD_TYPE_CHOICES),
        'weather_choices': dict(WEATHER_CONDITION_CHOICES),
        'stats': stats,
        'latest_occupancy': latest_occupancy,
    })


def visitor_flow_create(request):
    now = timezone.now()
    initial = {
        'record_date': now.date(),
        'record_time': now.time().replace(microsecond=0),
        'recorded_by': request.user.username if hasattr(request, 'user') and request.user.is_authenticated else '系统',
    }
    road_id = request.GET.get('road')
    if road_id:
        initial['road_section'] = road_id

    if request.method == 'POST':
        form = VisitorFlowRecordForm(request.POST)
        if form.is_valid():
            try:
                flow = form.save(commit=False)
                flow.full_clean()
                flow.save()
                messages.success(request, '客流记录登记成功。')
                return redirect('roads:visitor_flow_list')
            except ValidationError as e:
                for field, errors in e.message_dict.items():
                    for error in errors:
                        form.add_error(field, error)
    else:
        form = VisitorFlowRecordForm(initial=initial)
    return render(request, 'roads/visitor_flow_form.html', {
        'form': form, 'mode': 'create'
    })


def visitor_flow_edit(request, pk):
    flow = get_object_or_404(VisitorFlowRecord, pk=pk)
    if request.method == 'POST':
        form = VisitorFlowRecordForm(request.POST, instance=flow)
        if form.is_valid():
            try:
                f = form.save(commit=False)
                f.full_clean()
                f.save()
                messages.success(request, '客流记录更新成功。')
                return redirect('roads:visitor_flow_list')
            except ValidationError as e:
                for field, errors in e.message_dict.items():
                    for error in errors:
                        form.add_error(field, error)
    else:
        form = VisitorFlowRecordForm(instance=flow)
    return render(request, 'roads/visitor_flow_form.html', {
        'form': form, 'flow': flow, 'mode': 'edit'
    })


def visitor_flow_delete(request, pk):
    flow = get_object_or_404(VisitorFlowRecord, pk=pk)
    if request.method == 'POST':
        flow.delete()
        messages.success(request, '客流记录已删除。')
        return redirect('roads:visitor_flow_list')
    return render(request, 'roads/visitor_flow_confirm_delete.html', {
        'flow': flow
    })


def open_status_dashboard(request):
    roads = RoadSection.objects.all()
    now = timezone.now()

    road_statuses = []
    for road in roads:
        schedule = OpenSchedule.objects.filter(
            road_section=road, is_active=True
        ).first()
        active_controls = TemporaryControl.objects.filter(
            road_section=road, is_active=True,
            start_time__lte=now, end_time__gte=now
        )
        passage = RoadPassageStatus.objects.filter(
            road_section=road
        ).first()
        latest_flow = VisitorFlowRecord.objects.filter(
            road_section=road
        ).order_by('-record_date', '-record_time').first()
        suggestions = _generate_open_suggestion(road)

        current_status = 'closed'
        if schedule:
            current_status = schedule.get_current_open_status()
        if active_controls.exists():
            current_status = active_controls.first().open_status

        capacity_usage = 0
        max_capacity = 0
        if schedule:
            max_capacity = schedule.max_capacity
        if latest_flow and max_capacity > 0:
            capacity_usage = round(latest_flow.current_occupancy / max_capacity * 100, 1)

        road_statuses.append({
            'road': road,
            'schedule': schedule,
            'active_controls': active_controls,
            'passage': passage,
            'latest_flow': latest_flow,
            'suggestions': suggestions,
            'current_status': current_status,
            'current_status_display': dict(OPEN_STATUS_CHOICES).get(current_status, current_status),
            'capacity_usage': capacity_usage,
            'max_capacity': max_capacity,
        })

    total_open = sum(1 for rs in road_statuses if rs['current_status'] == 'open')
    total_partial = sum(1 for rs in road_statuses if rs['current_status'] == 'partial')
    total_restricted = sum(1 for rs in road_statuses if rs['current_status'] == 'restricted')
    total_closed = sum(1 for rs in road_statuses if rs['current_status'] == 'closed')

    stats = {
        'total_roads': len(road_statuses),
        'open': total_open,
        'partial': total_partial,
        'restricted': total_restricted,
        'closed': total_closed,
        'with_suggestions': sum(1 for rs in road_statuses if any(s['level'] != 'info' for s in rs['suggestions'])),
    }

    return render(request, 'roads/open_status_dashboard.html', {
        'road_statuses': road_statuses,
        'stats': stats,
        'open_status_choices': dict(OPEN_STATUS_CHOICES),
        'control_type_choices': dict(CONTROL_TYPE_CHOICES),
        'season_choices': dict(SEASON_CHOICES),
        'weather_choices': dict(WEATHER_CONDITION_CHOICES),
    })


def api_open_suggestion(request, road_pk):
    road = get_object_or_404(RoadSection, pk=road_pk)
    suggestions = _generate_open_suggestion(road)
    schedule = OpenSchedule.objects.filter(
        road_section=road, is_active=True
    ).first()

    effective_capacity = schedule.max_capacity if schedule else 0
    effective_open = str(schedule.open_time) if schedule else ''
    effective_close = str(schedule.close_time) if schedule else ''

    active_controls = TemporaryControl.objects.filter(
        road_section=road, is_active=True,
        start_time__lte=timezone.now(), end_time__gte=timezone.now()
    )
    if active_controls.exists():
        ctrl = active_controls.first()
        if ctrl.adjusted_capacity:
            effective_capacity = ctrl.adjusted_capacity
        if ctrl.adjusted_open_time:
            effective_open = str(ctrl.adjusted_open_time)
        if ctrl.adjusted_close_time:
            effective_close = str(ctrl.adjusted_close_time)

    return JsonResponse({
        'road_id': road.pk,
        'road_code': road.code,
        'road_name': road.name,
        'suggestions': suggestions,
        'effective_capacity': effective_capacity,
        'effective_open_time': effective_open,
        'effective_close_time': effective_close,
    })
